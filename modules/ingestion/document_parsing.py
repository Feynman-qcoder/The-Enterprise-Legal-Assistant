"""Unified, side-effect-free document parsing for supported knowledge files."""

from __future__ import annotations

import os
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

import yaml

from modules.ingestion.pdf_extract import extract_pages_pdf

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".txt", ".md", ".docx", ".doc"})
FRONT_MATTER_FIELDS = frozenset(
    {
        "doc_id",
        "title",
        "category",
        "document_type",
        "source_org",
        "source_url",
        "publish_date",
        "effective_date",
        "expiry_date",
        "status",
        "jurisdiction",
        "version",
    }
)


class DocumentParseError(ValueError):
    """Expected, user-facing parser failure with a stable error code."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class ParsedDocument:
    """Normalized parser output consumed by the shared chunking pipeline."""

    text: str
    metadata: dict[str, Any] = field(default_factory=dict)
    segments: tuple[str, ...] = ()


def _base_metadata(path: Path, *, mime_type: str) -> dict[str, Any]:
    return {
        "filename": path.name,
        "extension": path.suffix.lower(),
        "mime_type": mime_type,
        "title": path.stem,
        "source": path.name,
    }


def _read_utf8(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentParseError("INVALID_ENCODING", "文件必须使用 UTF-8 或 UTF-8-SIG 编码") from exc


def parse_pdf(path: Path) -> ParsedDocument:
    pages = tuple(page for page in extract_pages_pdf(path) if page.strip())
    text = "\n\n".join(pages).strip()
    if not text:
        raise DocumentParseError("EMPTY_DOCUMENT", "PDF 未提取到文本；扫描件需要 OCR")
    return ParsedDocument(
        text=text,
        segments=pages,
        metadata=_base_metadata(path, mime_type="application/pdf"),
    )


def parse_txt(path: Path) -> ParsedDocument:
    text = _read_utf8(path).strip()
    if not text:
        raise DocumentParseError("EMPTY_DOCUMENT", "TXT 文件没有可入库正文")
    return ParsedDocument(
        text=text,
        segments=(text,),
        metadata=_base_metadata(path, mime_type="text/plain"),
    )


def _split_front_matter(raw: str) -> tuple[dict[str, Any], str]:
    lines = raw.splitlines()
    if not lines or lines[0].strip() != "---":
        return {}, raw
    closing = next((i for i, line in enumerate(lines[1:], start=1) if line.strip() == "---"), None)
    if closing is None:
        raise DocumentParseError("INVALID_FRONT_MATTER", "Markdown Front Matter 缺少结束分隔符")
    try:
        loaded = yaml.safe_load("\n".join(lines[1:closing])) or {}
    except yaml.YAMLError as exc:
        raise DocumentParseError("INVALID_FRONT_MATTER", "Markdown Front Matter YAML 无法解析") from exc
    if not isinstance(loaded, dict):
        raise DocumentParseError("INVALID_FRONT_MATTER", "Markdown Front Matter 必须是键值映射")
    metadata = {str(k): v for k, v in loaded.items() if str(k) in FRONT_MATTER_FIELDS}
    return metadata, "\n".join(lines[closing + 1 :])


def parse_markdown(path: Path) -> ParsedDocument:
    raw = _read_utf8(path)
    front_matter, body = _split_front_matter(raw)
    body = body.strip()
    if not body:
        raise DocumentParseError("EMPTY_DOCUMENT", "Markdown 文件没有可入库正文")
    metadata = _base_metadata(path, mime_type="text/markdown")
    metadata.update(front_matter)
    return ParsedDocument(text=body, segments=(body,), metadata=metadata)


def _escape_table_cell(value: str) -> str:
    return " ".join(value.split()).replace("|", "\\|")


def _table_to_markdown(table: Any) -> str:
    rows = [[_escape_table_cell(cell.text) for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    header = normalized[0]
    output = [f"| {' | '.join(header)} |", f"| {' | '.join(['---'] * width)} |"]
    output.extend(f"| {' | '.join(row)} |" for row in normalized[1:])
    return "\n".join(output)


def parse_docx(path: Path) -> ParsedDocument:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise DocumentParseError("DEPENDENCY_MISSING", "DOCX 解析需要安装 python-docx") from exc

    try:
        document = Document(str(path))
    except Exception as exc:
        raise DocumentParseError("PARSE_FAILED", "DOCX 文件损坏或格式无效") from exc

    blocks: list[str] = []
    if hasattr(document, "iter_inner_content"):
        content = document.iter_inner_content()
    else:  # pragma: no cover - python-docx < 1.1 compatibility
        content = document.element.body.iterchildren()

    for item in content:
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if text:
                blocks.append(text)
        elif isinstance(item, Table):
            rendered = _table_to_markdown(item)
            if rendered:
                blocks.append(rendered)

    text = "\n\n".join(blocks).strip()
    if not text:
        raise DocumentParseError("EMPTY_DOCUMENT", "DOCX 文件没有可入库正文")
    metadata = _base_metadata(
        path,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    if document.core_properties.title:
        metadata["title"] = document.core_properties.title
    return ParsedDocument(text=text, segments=(text,), metadata=metadata)


def find_libreoffice() -> Path | None:
    """Locate LibreOffice from explicit config, PATH, then common install paths."""
    candidates = [os.getenv("LIBREOFFICE_PATH"), shutil.which("soffice"), shutil.which("soffice.exe")]
    if os.name == "nt":
        candidates.extend(
            [
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            ]
        )
    for candidate in candidates:
        if candidate:
            path = Path(candidate)
            if path.is_file():
                return path
    return None


def parse_doc(path: Path, *, timeout_seconds: int = 60) -> ParsedDocument:
    soffice = find_libreoffice()
    if soffice is None:
        raise DocumentParseError("PARSER_UNAVAILABLE", "Legacy .doc parsing requires LibreOffice")
    with TemporaryDirectory(prefix="xiaoyi-doc-") as temp_dir:
        output_dir = Path(temp_dir)
        try:
            completed = subprocess.run(
                [
                    str(soffice),
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    str(output_dir),
                    str(path),
                ],
                check=False,
                capture_output=True,
                text=True,
                shell=False,
                timeout=timeout_seconds,
            )
        except subprocess.TimeoutExpired as exc:
            raise DocumentParseError("CONVERSION_TIMEOUT", "LibreOffice 转换 .doc 超时") from exc
        converted = output_dir / f"{path.stem}.docx"
        if completed.returncode != 0 or not converted.is_file():
            raise DocumentParseError("CONVERSION_FAILED", "LibreOffice 无法转换该 .doc 文件")
        parsed = parse_docx(converted)
        metadata = dict(parsed.metadata)
        metadata.update(_base_metadata(path, mime_type="application/msword"))
        return ParsedDocument(text=parsed.text, segments=parsed.segments, metadata=metadata)


def parse_document(path: Path) -> ParsedDocument:
    """Dispatch a local file to the parser selected by its normalized suffix."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError("UNSUPPORTED_FILE_TYPE", f"不支持的文件类型：{suffix or '<none>'}")
    if not path.is_file():
        raise DocumentParseError("FILE_NOT_FOUND", "待解析文件不存在")
    parser = {
        ".pdf": parse_pdf,
        ".txt": parse_txt,
        ".md": parse_markdown,
        ".docx": parse_docx,
        ".doc": parse_doc,
    }[suffix]
    return parser(path)


def parser_available(extension: str) -> bool:
    suffix = extension.lower()
    if not suffix.startswith("."):
        suffix = f".{suffix}"
    return suffix in SUPPORTED_EXTENSIONS and (suffix != ".doc" or find_libreoffice() is not None)


def get_parser_capabilities() -> dict[str, dict[str, str | bool | None]]:
    capabilities: dict[str, dict[str, str | bool | None]] = {}
    for extension in sorted(SUPPORTED_EXTENSIONS):
        available = parser_available(extension)
        capabilities[extension] = {
            "available": available,
            "reason": None if available else "LibreOffice unavailable",
        }
    return capabilities
