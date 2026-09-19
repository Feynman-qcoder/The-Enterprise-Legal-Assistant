from pathlib import Path

from docx import Document

from modules.ingestion.document_parsing import parse_document


def _add_table(document: Document, label: str) -> None:
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = f"{label}金额"
    table.cell(0, 1).text = "审批人"
    table.cell(1, 0).text = ">500万"
    table.cell(1, 1).text = "总经理"


def test_docx_preserves_paragraph_table_order(tmp_path: Path) -> None:
    path = tmp_path / "policy.docx"
    document = Document()
    document.core_properties.title = "华辰采购制度"
    document.add_heading("标题", level=1)
    document.add_paragraph("正文A")
    _add_table(document, "表格A")
    document.add_paragraph("正文B")
    _add_table(document, "表格B")
    document.add_paragraph("正文C")
    document.save(path)

    parsed = parse_document(path)

    markers = ["标题", "正文A", "表格A金额", "正文B", "表格B金额", "正文C"]
    positions = [parsed.text.index(marker) for marker in markers]
    assert positions == sorted(positions)
    assert "| 表格A金额 | 审批人 |" in parsed.text
    assert parsed.metadata["title"] == "华辰采购制度"
