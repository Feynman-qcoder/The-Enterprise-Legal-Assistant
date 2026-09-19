"""Manual real-storage smoke for the MVP ingestion APIs (not collected by pytest)."""

from __future__ import annotations

import asyncio
import json
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from pymilvus import Collection
from sqlalchemy import select

from backend.app.main import app
from modules.database.models import FaqTab, LegalTab
from modules.database.session import get_async_engine, get_session_factory
from modules.milvus_store.client import ensure_milvus
from modules.milvus_store.collections import COLLECTION_FAQ, COLLECTION_LEGAL_CHILD

FIXTURE_DIR = Path(__file__).with_name("fixtures")


def _write_simple_pdf(path: Path, text: str) -> None:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    data.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    data.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    path.write_bytes(data)


def _prepare_fixtures() -> dict[str, Path]:
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    pdf = FIXTURE_DIR / "mvp_test.pdf"
    txt = FIXTURE_DIR / "mvp_test.txt"
    md = FIXTURE_DIR / "mvp_test.md"
    docx = FIXTURE_DIR / "mvp_test.docx"
    broken = FIXTURE_DIR / "broken.docx"
    _write_simple_pdf(pdf, "LEGAL_PDF_TEST_001 The designated approval officer is Chief Auditor.")
    txt.write_text("LEGAL_TXT_TEST_001 采购合同归档责任人是档案主管。", encoding="utf-8")
    md.write_text("# 紧急采购\n\nLEGAL_MD_TEST_001 紧急采购必须在三个工作日内补办审批。", encoding="utf-8")
    document = Document()
    document.add_heading("采购合同审批制度", level=1)
    document.add_paragraph("LEGAL_DOCX_TEST_001 超过500万元采购合同需要总经理审批。")
    document.save(docx)
    broken.write_bytes(b"broken docx")
    return {"pdf": pdf, "txt": txt, "md": md, "docx": docx, "broken": broken}


def _answer_from_sse(body: str) -> tuple[str, bool, bool]:
    chunks: list[str] = []
    done = False
    error = False
    for line in body.splitlines():
        if not line.startswith("data: "):
            continue
        payload = line[6:].strip()
        if payload == "[DONE]":
            done = True
            continue
        parsed = json.loads(payload)
        error = error or "error" in parsed
        if parsed.get("chunk"):
            chunks.append(str(parsed["chunk"]))
    return "".join(chunks), done, error


async def _mysql_summary(parent_ids: list[int], child_ids: list[int], faq_ids: list[int]) -> dict:
    factory = get_session_factory()
    async with factory() as session:
        parent_result = await session.execute(select(LegalTab).where(LegalTab.id.in_(parent_ids)))
        parents = list(parent_result.scalars().all())
        child_result = await session.execute(select(LegalTab).where(LegalTab.id.in_(child_ids)))
        children = list(child_result.scalars().all())
        faq_result = await session.execute(select(FaqTab).where(FaqTab.id.in_(faq_ids)))
        faq_rows = list(faq_result.scalars().all())
    await get_async_engine().dispose()
    return {
        "legal_parents": len(parents),
        "legal_children": len(children),
        "faq_rows": len(faq_rows),
        "parent_ids": [int(row.id) for row in parents],
        "child_parent_ids": [int(row.parent_id) for row in children if row.parent_id is not None],
    }


def main() -> None:
    fixtures = _prepare_fixtures()
    faq_fixture = Path(__file__).resolve().parents[1] / "document_parsing" / "fixtures" / "faq_test.xlsx"
    legacy_doc = (
        Path(__file__).resolve().parents[2]
        / "项目用到的软件及工具"
        / "mysql软件"
        / "mac系统"
        / "mysql安装步骤(mac版本).doc"
    )
    with TestClient(app) as client:
        with fixtures["pdf"].open("rb") as pdf, fixtures["txt"].open("rb") as txt, fixtures["md"].open(
            "rb"
        ) as md, fixtures["docx"].open("rb") as docx, fixtures["broken"].open("rb") as broken, legacy_doc.open(
            "rb"
        ) as legacy:
            document_response = client.post(
                "/api/knowledge/documents/upload",
                files=[
                    ("files", (fixtures["pdf"].name, pdf, "application/pdf")),
                    ("files", (fixtures["txt"].name, txt, "text/plain")),
                    ("files", (fixtures["md"].name, md, "text/markdown")),
                    (
                        "files",
                        (
                            fixtures["docx"].name,
                            docx,
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        ),
                    ),
                    ("files", (fixtures["broken"].name, broken, "application/octet-stream")),
                    ("files", (legacy_doc.name, legacy, "application/msword")),
                ],
            )
        with faq_fixture.open("rb") as faq:
            faq_response = client.post(
                "/api/knowledge/faqs/upload",
                files=[
                    (
                        "files",
                        (
                            faq_fixture.name,
                            faq,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        ),
                    )
                ],
            )

        questions = {
            "pdf": ("企业合同制度 LEGAL_PDF_TEST_001 指定的审批人是谁？", ("Chief Auditor", "审计")),
            "txt": ("LEGAL_TXT_TEST_001 采购合同归档责任人是谁？", ("档案主管",)),
            "md": ("LEGAL_MD_TEST_001 紧急采购应在多久内补办审批？", ("三个工作日",)),
            "docx": ("LEGAL_DOCX_TEST_001 超过500万元采购合同需要谁审批？", ("总经理",)),
            "faq": ("FAQ_TEST_001 什么情况下采购合同需要法务审核？", ("一百万元",)),
        }
        rag: dict[str, dict] = {}
        for kind, (question, expected_values) in questions.items():
            response = client.post("/api/chat/stream", json={"message": question})
            answer, done, error = _answer_from_sse(response.text)
            rag[kind] = {
                "http": response.status_code,
                "done": done,
                "error": error,
                "contains_expected": any(expected in answer for expected in expected_values),
                "answer_length": len(answer),
            }

    documents = document_response.json()
    faqs = faq_response.json()
    successful_documents = [item for item in documents["files"] if item["status"] == "success"]
    faq_ids = [value for item in faqs["files"] for value in item["record_ids"]]
    parent_ids = [value for item in successful_documents for value in item["parent_ids"]]
    legal_ids = [value for item in successful_documents for value in item["record_ids"]]
    mysql = asyncio.run(_mysql_summary(parent_ids, legal_ids, faq_ids))

    ensure_milvus()
    legal_collection = Collection(COLLECTION_LEGAL_CHILD)
    legal_collection.load()
    legal_entities = legal_collection.query(
        expr=f"id in [{','.join(str(value) for value in legal_ids)}]",
        output_fields=["id", "parent_id", "source_file"],
    )
    faq_collection = Collection(COLLECTION_FAQ)
    faq_collection.load()
    faq_entities = faq_collection.query(
        expr=f"id in [{','.join(str(value) for value in faq_ids)}]",
        output_fields=["id", "question"],
    )
    parent_mapping_ok = sorted(entity["parent_id"] for entity in legal_entities) == sorted(mysql["child_parent_ids"])
    print(
        json.dumps(
            {
                "documents": documents,
                "faqs": faqs,
                "mysql": mysql,
                "milvus": {
                    "legal_entities": len(legal_entities),
                    "faq_entities": len(faq_entities),
                    "parent_mapping_ok": parent_mapping_ok,
                },
                "rag": rag,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
