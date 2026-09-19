from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document

from offline.scripts import run_v2_file_ingest as ingest
from offline.scripts.run_v2_doc_sync import parse_package


SUBSTANTIVE_TEXT = (
    "第一条 为规范合同履行、保护各方合法权益，根据有关法律法规，订立本合同。\n\n"
    "第二条 各方应当遵循诚实信用原则，完整履行通知、协作、保密和数据安全义务。\n\n"
    "第三条 因履行本合同发生争议的，各方应当先行协商；协商不成的依法解决。"
)


@dataclass(frozen=True)
class Layout:
    corpus: Path
    baseline: Path
    live: Path
    base_file: Path
    base_entry: dict


def _write_manifest(path: Path, entries: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "".join(json.dumps(entry, ensure_ascii=False) + "\n" for entry in entries),
        encoding="utf-8",
    )


@pytest.fixture
def layout(tmp_path: Path) -> Layout:
    corpus = tmp_path / "data_corpus"
    corpus.mkdir()
    base_file = corpus / "BASE_001.md"
    base_file.write_text(f"# 测试合同\n\n{SUBSTANTIVE_TEXT}\n", encoding="utf-8")
    base_sha = ingest.compute_normalized_sha256(base_file)
    base_entry = {
        "logical_document_id": "BASE_001",
        "title": "测试合同",
        "canonical_file": base_file.name,
        "format": "md",
        "source_org": "测试机构",
        "cleaning_status": "PASS",
        "canonical": True,
        "normalized_sha256": base_sha,
    }
    baseline = tmp_path / "frozen" / "manifest.jsonl"
    live = tmp_path / "data" / "manifests" / "live.jsonl"
    _write_manifest(baseline, [base_entry])
    return Layout(corpus, baseline, live, base_file, base_entry)


def _common_args(layout: Layout) -> list[str]:
    return [
        "--manifest",
        str(layout.live),
        "--baseline-manifest",
        str(layout.baseline),
        "--corpus-dir",
        str(layout.corpus),
    ]


def _fake_package(entry: dict, *, child_id: str = "a" * 32):
    identity = f"{entry['logical_document_id']}@sha256:{entry['normalized_sha256']}"
    raw = {
        "document_identity": identity,
        "source_file": entry["canonical_file"],
        "identity_sha256": entry["normalized_sha256"],
        "identity_sha256_kind": "normalized_sha256",
        "parents": [
            {
                "section_path": ["第一条"],
                "content": SUBSTANTIVE_TEXT,
                "metadata": {},
                "children": [
                    {
                        "content": SUBSTANTIVE_TEXT,
                        "section_path": ["第一条"],
                        "metadata": {"chunk_index": 0},
                        "child_id": child_id,
                    }
                ],
            }
        ],
    }
    return raw, parse_package(raw)


def _install_fake_builder(monkeypatch: pytest.MonkeyPatch) -> None:
    def fake_builder(**kwargs):
        return _fake_package(kwargs["entry"])

    monkeypatch.setattr(ingest, "run_package_builder", fake_builder)


@pytest.mark.asyncio
async def test_existing_manifest_document_dry_run(layout: Layout) -> None:
    package_one = layout.live.parent / "retry-one.json"
    package_two = layout.live.parent / "retry-two.json"
    report = await ingest.run(
        [
            "--doc",
            "BASE_001",
            "--dry-run",
            "--package-out",
            str(package_one),
            *_common_args(layout),
        ]
    )
    await ingest.run(
        [
            "--doc",
            "BASE_001",
            "--dry-run",
            "--package-out",
            str(package_two),
            *_common_args(layout),
        ]
    )
    first = json.loads(package_one.read_text(encoding="utf-8"))
    second = json.loads(package_two.read_text(encoding="utf-8"))
    first_ids = [child["child_id"] for parent in first["parents"] for child in parent["children"]]
    second_ids = [child["child_id"] for parent in second["parents"] for child in parent["children"]]
    assert report["ingestion_status"] == "DRY_RUN"
    assert report["parent_count"] >= 1
    assert report["child_count"] >= 1
    assert report["normalized_sha256"] == layout.base_entry["normalized_sha256"]
    assert first_ids == second_ids
    assert len(ingest.load_manifest(layout.live)) == 1


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix", [".md", ".docx"])
async def test_new_md_and_docx_dry_run_do_not_update_manifest(
    layout: Layout,
    suffix: str,
) -> None:
    source = layout.corpus / f"NEW_001{suffix}"
    if suffix == ".md":
        source.write_text(f"# 新合同\n\n{SUBSTANTIVE_TEXT}\n", encoding="utf-8")
    else:
        document = Document()
        document.add_heading("新合同", level=1)
        document.add_paragraph(SUBSTANTIVE_TEXT)
        document.save(source)
    package_out = layout.live.parent / f"new-{suffix[1:]}.json"
    report = await ingest.run(
        [
            "--file",
            str(source),
            "--logical-document-id",
            f"NEW_{suffix[1:].upper()}",
            "--title",
            "新合同",
            "--source-org",
            "来源机构",
            "--package-out",
            str(package_out),
            *_common_args(layout),
        ]
    )
    package = json.loads(package_out.read_text(encoding="utf-8"))
    child_ids = [
        child["child_id"] for parent in package["parents"] for child in parent["children"]
    ]
    assert report["ingestion_status"] == "DRY_RUN"
    assert package["identity_sha256_kind"] == "normalized_sha256"
    assert child_ids and len(child_ids) == len(set(child_ids))
    assert len(ingest.load_manifest(layout.live)) == 1


@pytest.mark.asyncio
async def test_manifest_missing_fails_fast(tmp_path: Path) -> None:
    corpus = tmp_path / "data_corpus"
    corpus.mkdir()
    with pytest.raises(ingest.IngestionFailure, match="Manifest not found") as caught:
        await ingest.run(
            [
                "--doc",
                "MISSING",
                "--manifest",
                str(tmp_path / "live.jsonl"),
                "--baseline-manifest",
                str(tmp_path / "missing.jsonl"),
                "--corpus-dir",
                str(corpus),
            ]
        )
    assert caught.value.stage == "manifest"


@pytest.mark.asyncio
async def test_cleaner_failure_stops_before_builder(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        ingest,
        "compute_normalized_sha256",
        lambda _path: (_ for _ in ()).throw(ingest.IngestionFailure("cleaner", "boom")),
    )
    monkeypatch.setattr(
        ingest,
        "run_package_builder",
        lambda **_kwargs: (_ for _ in ()).throw(AssertionError("builder must not run")),
    )
    with pytest.raises(ingest.IngestionFailure, match="boom") as caught:
        await ingest.run(["--doc", "BASE_001", *_common_args(layout)])
    assert caught.value.stage == "cleaner"


@pytest.mark.asyncio
async def test_source_change_during_package_build_fails_before_storage(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = layout.corpus / "CHANGING.md"
    source.write_text(SUBSTANTIVE_TEXT, encoding="utf-8")
    monkeypatch.setattr(ingest, "compute_normalized_sha256", lambda _path: "c" * 64)

    def changing_builder(**kwargs):
        raw, package = _fake_package(kwargs["entry"])
        source.write_text(SUBSTANTIVE_TEXT + "\n第四条 文件已变化。", encoding="utf-8")
        return raw, package

    monkeypatch.setattr(ingest, "run_package_builder", changing_builder)
    monkeypatch.setattr(
        ingest,
        "run_apply_preflight",
        lambda **_kwargs: (_ for _ in ()).throw(AssertionError("storage must not run")),
    )
    with pytest.raises(ingest.IngestionFailure, match="source file changed") as caught:
        await ingest.run(
            [
                "--file",
                str(source),
                "--logical-document-id",
                "CHANGING",
                "--title",
                "变化文档",
                *_common_args(layout),
            ]
        )
    assert caught.value.stage == "source_changed"
    assert [entry["logical_document_id"] for entry in ingest.load_manifest(layout.live)] == [
        "BASE_001"
    ]


def test_manifest_lock_rejects_concurrent_ingestion(tmp_path: Path) -> None:
    manifest = tmp_path / "manifests" / "live.jsonl"
    with ingest.manifest_run_lock(manifest):
        with pytest.raises(ingest.IngestionFailure, match="another V2 ingestion") as caught:
            with ingest.manifest_run_lock(manifest):
                pass
    assert caught.value.stage == "manifest_lock"


def test_logical_document_id_respects_mysql_identity_limit(layout: Layout) -> None:
    valid = dict(layout.base_entry, logical_document_id="A" * 119)
    assert ingest._validate_manifest_entry(valid, line_no=1)
    invalid = dict(layout.base_entry, logical_document_id="A" * 120)
    with pytest.raises(ingest.IngestionFailure, match="unsafe logical_document_id"):
        ingest._validate_manifest_entry(invalid, line_no=1)


def test_invalid_cli_returns_staged_nonzero_without_pass(capsys: pytest.CaptureFixture) -> None:
    exit_code = ingest.main(["--doc", "BASE_001", "--dry-run", "--apply"])
    captured = capsys.readouterr()
    assert exit_code == 2
    assert "FAILED_STAGE=cli" in captured.err
    assert "INGESTION_STATUS=FAIL" in captured.err
    assert "INGESTION_STATUS=PASS" not in captured.out + captured.err


@pytest.mark.parametrize(
    ("returncode", "stderr", "expected_stage"),
    [
        (4, "[FAIL] stage=cleaner __CLEAN_FAIL_RuntimeError", "cleaner"),
        (3, "[FAIL] Chunk Contract V1 校验失败", "chunk_contract"),
    ],
)
def test_builder_cleaner_and_contract_failures_are_fail_fast(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    returncode: int,
    stderr: str,
    expected_stage: str,
) -> None:
    monkeypatch.setattr(
        ingest.subprocess,
        "run",
        lambda *_args, **_kwargs: SimpleNamespace(
            returncode=returncode,
            stderr=stderr,
            stdout="",
        ),
    )
    entry = {
        "logical_document_id": "TEST_001",
        "canonical_file": "TEST_001.md",
        "normalized_sha256": "a" * 64,
    }
    with pytest.raises(ingest.IngestionFailure) as caught:
        ingest.run_package_builder(
            entry=entry,
            manifest_path=tmp_path / "manifest.jsonl",
            corpus_dir=tmp_path,
        )
    assert caught.value.stage == expected_stage


@pytest.mark.asyncio
@pytest.mark.parametrize("stage", ["mysql_preflight", "milvus_preflight"])
async def test_unavailable_storage_fails_before_write(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
    stage: str,
) -> None:
    _install_fake_builder(monkeypatch)
    monkeypatch.setattr(
        ingest,
        "compute_normalized_sha256",
        lambda _path: layout.base_entry["normalized_sha256"],
    )

    async def failing_preflight(**_kwargs):
        raise ingest.IngestionFailure(stage, "unavailable")

    async def must_not_sync(*_args, **_kwargs):
        raise AssertionError("write must not run after failed preflight")

    monkeypatch.setattr(ingest, "run_apply_preflight", failing_preflight)
    monkeypatch.setattr(ingest, "sync_package", must_not_sync)
    with pytest.raises(ingest.IngestionFailure) as caught:
        await ingest.run(["--doc", "BASE_001", "--apply", *_common_args(layout)])
    assert caught.value.stage == stage


@pytest.mark.asyncio
async def test_apply_authorization_defaults_to_offline_dry_run(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _install_fake_builder(monkeypatch)
    monkeypatch.setattr(
        ingest,
        "compute_normalized_sha256",
        lambda _path: layout.base_entry["normalized_sha256"],
    )

    async def must_not_preflight(**_kwargs):
        raise AssertionError("default dry-run must not touch MySQL or Milvus")

    monkeypatch.setattr(ingest, "run_apply_preflight", must_not_preflight)
    report = await ingest.run(["--doc", "BASE_001", *_common_args(layout)])
    assert report["ingestion_status"] == "DRY_RUN"
    assert report["database_writes"] == 0
    assert report["milvus_writes"] == 0


@pytest.mark.asyncio
async def test_duplicate_logical_document_is_rejected(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        ingest,
        "compute_normalized_sha256",
        lambda _path: layout.base_entry["normalized_sha256"],
    )
    with pytest.raises(ingest.IngestionFailure) as caught:
        await ingest.run(
            [
                "--file",
                str(layout.base_file),
                "--logical-document-id",
                "BASE_001",
                "--title",
                "重复合同",
                *_common_args(layout),
            ]
        )
    assert caught.value.stage == "duplicate_document"
    assert len(ingest.load_manifest(layout.live)) == 1


@pytest.mark.asyncio
async def test_duplicate_canonical_file_with_new_logical_id_is_rejected(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        ingest,
        "compute_normalized_sha256",
        lambda _path: layout.base_entry["normalized_sha256"],
    )
    with pytest.raises(ingest.IngestionFailure) as caught:
        await ingest.run(
            [
                "--file",
                str(layout.base_file),
                "--logical-document-id",
                "OTHER_001",
                "--title",
                "文件名冲突",
                *_common_args(layout),
            ]
        )
    assert caught.value.stage == "manifest_conflict"


@pytest.mark.asyncio
async def test_changed_sha_requires_explicit_replace(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(ingest, "compute_normalized_sha256", lambda _path: "c" * 64)
    with pytest.raises(ingest.IngestionFailure, match="--replace-existing") as caught:
        await ingest.run(["--doc", "BASE_001", *_common_args(layout)])
    assert caught.value.stage == "document_identity"
    assert ingest.load_manifest(layout.live)[0] == layout.base_entry


def test_store_identity_conflict_requires_replace() -> None:
    old_identity = f"BASE_001@sha256:{'a' * 64}"
    new_identity = f"BASE_001@sha256:{'b' * 64}"
    preflight = ingest.ApplyPreflight(
        version_before=1,
        mysql_identities=frozenset({old_identity}),
        milvus_identities=frozenset({old_identity}),
        legacy_count_before=1,
        vector_dimension=1024,
        embedding_service=object(),
    )
    with pytest.raises(ingest.IngestionFailure, match="Old identities"):
        ingest.determine_old_identities(
            new_identity=new_identity,
            preflight=preflight,
            replace_existing=False,
        )
    assert ingest.determine_old_identities(
        new_identity=new_identity,
        preflight=preflight,
        replace_existing=True,
    ) == [old_identity]


def test_cross_store_identity_disagreement_fails_closed() -> None:
    identity = f"BASE_001@sha256:{'a' * 64}"
    preflight = ingest.ApplyPreflight(
        version_before=1,
        mysql_identities=frozenset({identity}),
        milvus_identities=frozenset(),
        legacy_count_before=1,
        vector_dimension=1024,
        embedding_service=object(),
    )
    with pytest.raises(ingest.IngestionFailure, match="sets disagree"):
        ingest.determine_old_identities(
            new_identity=identity,
            preflight=preflight,
            replace_existing=True,
        )


@pytest.mark.asyncio
async def test_apply_reports_verified_counts_and_preserves_legacy_collection(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    new_file = layout.corpus / "NEW_APPLY.md"
    new_file.write_text(SUBSTANTIVE_TEXT, encoding="utf-8")
    new_sha = "d" * 64
    monkeypatch.setattr(ingest, "compute_normalized_sha256", lambda _path: new_sha)
    _install_fake_builder(monkeypatch)
    preflight = ingest.ApplyPreflight(
        version_before=10,
        mysql_identities=frozenset(),
        milvus_identities=frozenset(),
        legacy_count_before=77,
        vector_dimension=1024,
        embedding_service=object(),
    )

    async def fake_preflight(**_kwargs):
        return preflight

    async def fake_sync(package, collection_name, *, embedding_service):
        assert collection_name == ingest.V2_COLLECTION
        assert embedding_service is preflight.embedding_service
        return {
            "insert": {
                "mysql_parent_inserted": len(package.parents),
                "mysql_map_inserted": 1,
                "milvus_inserted": 1,
            },
            "new_corpus_version": 11,
        }

    async def fake_verify(**_kwargs):
        return ingest.WriteVerification(
            parent_count=1,
            mapping_count=1,
            milvus_count=1,
            orphan_count=0,
            version_after=11,
            legacy_count_after=77,
        )

    monkeypatch.setattr(ingest, "run_apply_preflight", fake_preflight)
    monkeypatch.setattr(ingest, "sync_package", fake_sync)
    monkeypatch.setattr(ingest, "verify_write", fake_verify)
    report = await ingest.run(
        [
            "--file",
            str(new_file),
            "--logical-document-id",
            "NEW_APPLY",
            "--title",
            "新增合同",
            "--apply",
            *_common_args(layout),
        ]
    )
    live_entries = ingest.load_manifest(layout.live)
    assert report["ingestion_status"] == "PASS"
    assert report["mysql_parent_written"] == 1
    assert report["mysql_mapping_written"] == 1
    assert report["milvus_written"] == 1
    assert report["orphan_count"] == 0
    assert report["vector_dimension"] == 1024
    assert report["corpus_version_before"] == 10
    assert report["corpus_version_after"] == 11
    assert report["legacy_count_before"] == report["legacy_count_after"] == 77
    assert {entry["logical_document_id"] for entry in live_entries} == {
        "BASE_001",
        "NEW_APPLY",
    }


@pytest.mark.asyncio
async def test_replace_existing_orders_sync_delete_verify_then_manifest_commit(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    new_sha = "9" * 64
    old_identity = ingest._manifest_identity(layout.base_entry)
    new_identity = f"BASE_001@sha256:{new_sha}"
    monkeypatch.setattr(ingest, "compute_normalized_sha256", lambda _path: new_sha)
    _install_fake_builder(monkeypatch)
    preflight = ingest.ApplyPreflight(
        version_before=10,
        mysql_identities=frozenset({old_identity}),
        milvus_identities=frozenset({old_identity}),
        legacy_count_before=77,
        vector_dimension=1024,
        embedding_service=object(),
    )
    events: list[str] = []

    async def fake_preflight(**_kwargs):
        events.append("preflight")
        return preflight

    async def fake_sync(package, collection_name, *, embedding_service):
        assert package.document_identity == new_identity
        assert collection_name == ingest.V2_COLLECTION
        assert embedding_service is preflight.embedding_service
        events.append("sync_new")
        return {
            "insert": {
                "mysql_parent_inserted": 1,
                "mysql_map_inserted": 1,
                "milvus_inserted": 1,
            }
        }

    async def fake_delete(identity, collection_name):
        assert identity == old_identity
        assert collection_name == ingest.V2_COLLECTION
        events.append("delete_old")
        return {}

    async def fake_verify(*, package, old_identities, preflight):
        assert package.document_identity == new_identity
        assert old_identities == [old_identity]
        assert preflight is not None
        events.append("verify")
        assert ingest.load_manifest(layout.live)[0]["normalized_sha256"] != new_sha
        return ingest.WriteVerification(1, 1, 1, 0, 12, 77)

    monkeypatch.setattr(ingest, "run_apply_preflight", fake_preflight)
    monkeypatch.setattr(ingest, "sync_package", fake_sync)
    monkeypatch.setattr(ingest, "sync_delete", fake_delete)
    monkeypatch.setattr(ingest, "verify_write", fake_verify)
    report = await ingest.run(
        [
            "--doc",
            "BASE_001",
            "--replace-existing",
            "--apply",
            *_common_args(layout),
        ]
    )
    assert events == ["preflight", "sync_new", "delete_old", "verify"]
    assert ingest.load_manifest(layout.live)[0]["normalized_sha256"] == new_sha
    assert report["old_identities_removed"] == [old_identity]


@pytest.mark.asyncio
async def test_failed_write_never_commits_candidate_manifest(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    new_file = layout.corpus / "NEW_FAIL.md"
    new_file.write_text(SUBSTANTIVE_TEXT, encoding="utf-8")
    monkeypatch.setattr(ingest, "compute_normalized_sha256", lambda _path: "e" * 64)
    _install_fake_builder(monkeypatch)
    preflight = ingest.ApplyPreflight(
        version_before=10,
        mysql_identities=frozenset(),
        milvus_identities=frozenset(),
        legacy_count_before=77,
        vector_dimension=1024,
        embedding_service=object(),
    )

    async def fake_preflight(**_kwargs):
        return preflight

    async def failing_sync(*_args, **_kwargs):
        raise RuntimeError("write failed")

    monkeypatch.setattr(ingest, "run_apply_preflight", fake_preflight)
    monkeypatch.setattr(ingest, "sync_package", failing_sync)
    with pytest.raises(ingest.IngestionFailure) as caught:
        await ingest.run(
            [
                "--file",
                str(new_file),
                "--logical-document-id",
                "NEW_FAIL",
                "--title",
                "失败合同",
                "--apply",
                *_common_args(layout),
            ]
        )
    assert caught.value.stage == "write"
    assert [entry["logical_document_id"] for entry in ingest.load_manifest(layout.live)] == [
        "BASE_001"
    ]


@pytest.mark.asyncio
async def test_failed_post_verification_never_commits_candidate_manifest(
    layout: Layout,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    new_file = layout.corpus / "NEW_VERIFY_FAIL.md"
    new_file.write_text(SUBSTANTIVE_TEXT, encoding="utf-8")
    monkeypatch.setattr(ingest, "compute_normalized_sha256", lambda _path: "f" * 64)
    _install_fake_builder(monkeypatch)
    preflight = ingest.ApplyPreflight(
        version_before=10,
        mysql_identities=frozenset(),
        milvus_identities=frozenset(),
        legacy_count_before=77,
        vector_dimension=1024,
        embedding_service=object(),
    )

    async def fake_preflight(**_kwargs):
        return preflight

    async def fake_sync(*_args, **_kwargs):
        return {
            "insert": {
                "mysql_parent_inserted": 1,
                "mysql_map_inserted": 1,
                "milvus_inserted": 1,
            }
        }

    async def failing_verify(**_kwargs):
        raise ingest.IngestionFailure("post_verify", "orphan=1")

    monkeypatch.setattr(ingest, "run_apply_preflight", fake_preflight)
    monkeypatch.setattr(ingest, "sync_package", fake_sync)
    monkeypatch.setattr(ingest, "verify_write", failing_verify)
    with pytest.raises(ingest.IngestionFailure) as caught:
        await ingest.run(
            [
                "--file",
                str(new_file),
                "--logical-document-id",
                "NEW_VERIFY_FAIL",
                "--title",
                "核验失败合同",
                "--apply",
                *_common_args(layout),
            ]
        )
    assert caught.value.stage == "post_verify"
    assert [entry["logical_document_id"] for entry in ingest.load_manifest(layout.live)] == [
        "BASE_001"
    ]


def _install_verification_fakes(
    monkeypatch: pytest.MonkeyPatch,
    *,
    legacy_count: int = 77,
    orphan: bool = False,
) -> None:
    import pymilvus

    import modules.database.session as session_module

    parent_id = "p" * 64
    mapped_parent_id = "q" * 64 if orphan else parent_id
    child_id = "a" * 32

    class FakeResult:
        def __init__(self, rows):
            self._rows = rows

        def all(self):
            return list(self._rows)

        def first(self):
            return self._rows[0] if self._rows else None

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def execute(self, statement, _params=None):
            sql = str(statement)
            if "SELECT external_parent_id, content" in sql:
                return FakeResult(
                    [(parent_id, SUBSTANTIVE_TEXT, '["第一条"]', "{}")]
                )
            if "SELECT m.child_id" in sql:
                return FakeResult([(child_id, mapped_parent_id)])
            if "SELECT DISTINCT document_identity" in sql:
                return FakeResult([(f"BASE_001@sha256:{'b' * 64}",)])
            if "SELECT meta_value FROM corpus_meta" in sql:
                return FakeResult([("11",)])
            raise AssertionError(f"unexpected SQL: {sql}")

    class FakeFactory:
        def __call__(self):
            return FakeSession()

    class FakeCollection:
        def __init__(self, name):
            self.name = name
            self.num_entities = legacy_count if name == ingest.LEGACY_COLLECTION else 1

        def query(self, **kwargs):
            assert self.name == ingest.V2_COLLECTION
            if " like " in kwargs["expr"]:
                return [{"document_identity": f"BASE_001@sha256:{'b' * 64}"}]
            return [
                {
                    "child_id": child_id,
                    "external_parent_id": mapped_parent_id,
                    "document_identity": f"BASE_001@sha256:{'b' * 64}",
                    "source_file": "BASE_001.md",
                }
            ]

    class FakeUtility:
        @staticmethod
        def has_collection(name):
            return name == ingest.LEGACY_COLLECTION

    monkeypatch.setattr(session_module, "get_session_factory", lambda: FakeFactory())
    monkeypatch.setattr(pymilvus, "Collection", FakeCollection)
    monkeypatch.setattr(pymilvus, "utility", FakeUtility())


@pytest.mark.asyncio
async def test_write_verification_queries_fake_mysql_and_milvus(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _install_verification_fakes(monkeypatch)
    entry = {
        "logical_document_id": "BASE_001",
        "canonical_file": "BASE_001.md",
        "normalized_sha256": "b" * 64,
    }
    _raw, package = _fake_package(entry)
    preflight = ingest.ApplyPreflight(
        version_before=10,
        mysql_identities=frozenset(),
        milvus_identities=frozenset(),
        legacy_count_before=77,
        vector_dimension=1024,
        embedding_service=object(),
    )
    result = await ingest.verify_write(
        package=package,
        old_identities=[],
        preflight=preflight,
    )
    assert result == ingest.WriteVerification(1, 1, 1, 0, 11, 77)


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("legacy_count", "orphan", "message"),
    [(78, False, "Legacy collection changed"), (77, True, "orphan count")],
)
async def test_write_verification_fails_on_legacy_change_or_orphan(
    monkeypatch: pytest.MonkeyPatch,
    legacy_count: int,
    orphan: bool,
    message: str,
) -> None:
    _install_verification_fakes(
        monkeypatch,
        legacy_count=legacy_count,
        orphan=orphan,
    )
    entry = {
        "logical_document_id": "BASE_001",
        "canonical_file": "BASE_001.md",
        "normalized_sha256": "b" * 64,
    }
    _raw, package = _fake_package(entry)
    preflight = ingest.ApplyPreflight(
        version_before=10,
        mysql_identities=frozenset(),
        milvus_identities=frozenset(),
        legacy_count_before=77,
        vector_dimension=1024,
        embedding_service=object(),
    )
    with pytest.raises(ingest.IngestionFailure, match=message):
        await ingest.verify_write(
            package=package,
            old_identities=[],
            preflight=preflight,
        )
