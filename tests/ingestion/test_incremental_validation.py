from io import BytesIO

import pytest
from openpyxl import Workbook

from modules.ingestion.document_parsing import ParsedDocument
from modules.ingestion.incremental import ingest_faq_workbook, ingest_legal_document
import modules.ingestion.incremental as incremental


def _generic_xlsx() -> bytes:
    stream = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["合同金额", "审批人"])
    sheet.append([">500万", "总经理"])
    workbook.save(stream)
    return stream.getvalue()


@pytest.mark.asyncio
async def test_documents_reject_doc_and_generic_xlsx() -> None:
    for filename in ("legacy.doc", "approval_matrix.xlsx"):
        result = await ingest_legal_document(filename, b"not used")
        assert result.status == "failed"
        assert result.error_code == "UNSUPPORTED_FILE_TYPE"


@pytest.mark.asyncio
async def test_faq_endpoint_rejects_generic_xlsx_schema() -> None:
    result = await ingest_faq_workbook("approval_matrix.xlsx", _generic_xlsx())
    assert result.status == "failed"
    assert result.error_code == "FAQ_SCHEMA_INVALID"


@pytest.mark.asyncio
async def test_cleaning_quality_gate_fails_before_embedding_or_storage(monkeypatch) -> None:
    parsed = ParsedDocument(
        text="第一条 正文。",
        segments=("第一条 正文。",),
        metadata={"extension": ".md"},
    )
    reviewed = ParsedDocument(
        text=parsed.text,
        segments=parsed.segments,
        metadata={
            **parsed.metadata,
            "cleaning": {
                "quality_status": "NEEDS_REVIEW",
                "quality_reasons": ["residual_boilerplate_hits=1: jiathis.com"],
            },
        },
    )
    monkeypatch.setattr(incremental, "parse_document", lambda _path: parsed)
    monkeypatch.setattr(incremental, "clean_parsed_document", lambda _parsed: reviewed)

    class MustNotEmbed:
        def __init__(self) -> None:
            raise AssertionError("quality gate must run before embedding")

    monkeypatch.setattr(incremental, "LocalEmbeddingService", MustNotEmbed)
    monkeypatch.setattr(
        incremental,
        "get_session_factory",
        lambda: (_ for _ in ()).throw(AssertionError("quality gate must run before MySQL")),
    )
    monkeypatch.setattr(
        incremental,
        "insert_legal_incremental",
        lambda *_args: (_ for _ in ()).throw(AssertionError("quality gate must run before Milvus")),
    )

    result = await ingest_legal_document("review.md", b"placeholder")

    assert result.status == "failed"
    assert result.error_code == "CLEANING_QUALITY_GATE_FAILED"
